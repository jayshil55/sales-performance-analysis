from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import DataBarRule

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "Sales_Data.xlsx"
OUT_FILE = HERE / "Sales_Analysis_Dashboard.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="176B87")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=18, bold=True, color="176B87")
SUB_FONT = Font(name="Calibri", size=9, italic=True, color="555555")
KPI_LABEL_FONT = Font(name="Calibri", size=9, bold=True, color="444444")
KPI_VALUE_FONT = Font(name="Calibri", size=15, bold=True, color="176B87")
KPI_TEXT_FONT = Font(name="Calibri", size=12, bold=True, color="176B87")
KPI_FILL = PatternFill("solid", fgColor="F0F6F7")
KPI_BOX_BORDER = Border(
    left=Side(style="medium", color="176B87"),
    right=Side(style="medium", color="176B87"),
    top=Side(style="medium", color="176B87"),
    bottom=Side(style="medium", color="176B87"),
)
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="176B87")

raw = pd.read_excel(SOURCE, sheet_name="Sales Data")
raw.columns = raw.columns.str.strip()

sales = raw.drop_duplicates().copy()
sales["Purchase Date"] = pd.to_datetime(sales["Purchase Date"], errors="coerce")
for col in ["Quantity", "Unit Price", "Total Price", "Discount Applied (%)", "Discount Amount", "Final Price"]:
    sales[col] = pd.to_numeric(sales[col], errors="coerce")
sales = sales.dropna(subset=["Order ID", "Purchase Date", "Product ID", "Final Price"])

products = pd.read_excel(SOURCE, sheet_name="Products")
products.columns = products.columns.str.strip()

product_sales = sales.groupby("Product ID", as_index=False).agg(
    units_sold=("Quantity", "sum"),
    revenue=("Final Price", "sum"),
    gross_sales=("Total Price", "sum"),
    discount_amount=("Discount Amount", "sum"),
)
product_sales = product_sales.merge(products[["Product ID", "Product Name", "Category"]], on="Product ID", how="left")
product_sales = product_sales.sort_values("revenue", ascending=False).reset_index(drop=True)

category_sales = product_sales.groupby("Category", as_index=False).agg(
    units_sold=("units_sold", "sum"),
    revenue=("revenue", "sum"),
    gross_sales=("gross_sales", "sum"),
).sort_values("revenue", ascending=False).reset_index(drop=True)

location_sales = sales.groupby("Store Location", as_index=False).agg(
    orders=("Order ID", "nunique"),
    units_sold=("Quantity", "sum"),
    revenue=("Final Price", "sum"),
).sort_values("revenue", ascending=False).reset_index(drop=True)

monthly = sales.assign(Month=sales["Purchase Date"].dt.to_period("M").astype(str)).groupby("Month", as_index=False).agg(
    orders=("Order ID", "nunique"),
    units_sold=("Quantity", "sum"),
    revenue=("Final Price", "sum"),
    gross_sales=("Total Price", "sum"),
    discount_amount=("Discount Amount", "sum"),
)
monthly = monthly.sort_values("Month").reset_index(drop=True)

top_revenue_prod_name = product_sales.iloc[0]["Product Name"]

wb = Workbook()

ws_dash = wb.active
ws_dash.title = "Dashboard"

ws_raw = wb.create_sheet("Raw Data")
ws_clean = wb.create_sheet("Cleaned Data")
ws_analysis = wb.create_sheet("Analysis")

for col in range(1, len(raw.columns) + 2):
    ws_raw.column_dimensions[get_column_letter(col)].width = 16
    ws_clean.column_dimensions[get_column_letter(col)].width = 16

for row in dataframe_to_rows(raw, index=False, header=True):
    ws_raw.append(row)
for row in dataframe_to_rows(sales, index=False, header=True):
    ws_clean.append(row)

for ci in range(1, len(raw.columns) + 1):
    c = ws_raw.cell(row=1, column=ci)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")

for ci in range(1, len(sales.columns) + 1):
    c = ws_clean.cell(row=1, column=ci)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")

last_raw = ws_raw.max_row
last_clean = ws_clean.max_row

raw_table = Table(displayName="RawDataTable", ref=f"A1:{get_column_letter(len(raw.columns))}{last_raw}")
raw_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws_raw.add_table(raw_table)

clean_table = Table(displayName="CleanedDataTable", ref=f"A1:{get_column_letter(len(sales.columns))}{last_clean}")
clean_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws_clean.add_table(clean_table)

ws_raw.freeze_panes = "A2"
ws_clean.freeze_panes = "A2"

for r in range(2, last_clean + 1):
    ws_clean.cell(row=r, column=2).number_format = "yyyy-mm-dd"
    ws_clean.cell(row=r, column=7).number_format = "#,##0"
    ws_clean.cell(row=r, column=8).number_format = "#,##0.00"
    ws_clean.cell(row=r, column=9).number_format = "#,##0.00"
    ws_clean.cell(row=r, column=10).number_format = "0"
    ws_clean.cell(row=r, column=11).number_format = "#,##0.00"
    ws_clean.cell(row=r, column=12).number_format = "#,##0.00"

ws_analysis.sheet_view.showGridLines = True
ws_analysis["A1"] = "Sales Analysis Summaries"
ws_analysis["A1"].font = TITLE_FONT
ws_analysis.row_dimensions[1].height = 28

ws_analysis["A3"] = "Monthly Performance"
ws_analysis["A3"].font = SECTION_FONT
monthly_headers = ["Month", "Orders", "Units Sold", "Gross Sales", "Discounts", "Net Revenue"]
for ci, h in enumerate(monthly_headers, 1):
    c = ws_analysis.cell(row=4, column=ci, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")

for idx, r in monthly.iterrows():
    row_num = 5 + idx
    ws_analysis.cell(row=row_num, column=1, value=pd.to_datetime(r["Month"] + "-01").to_pydatetime()).number_format = "mmm-yy"
    ws_analysis.cell(row=row_num, column=2, value=int(r["orders"])).number_format = "#,##0"
    ws_analysis.cell(row=row_num, column=3, value=int(r["units_sold"])).number_format = "#,##0"
    ws_analysis.cell(row=row_num, column=4, value=float(r["gross_sales"])).number_format = "#,##0.00"
    ws_analysis.cell(row=row_num, column=5, value=float(r["discount_amount"])).number_format = "#,##0.00"
    ws_analysis.cell(row=row_num, column=6, value=float(r["revenue"])).number_format = "#,##0.00"

monthly_last_row = 4 + len(monthly)

ws_analysis.cell(row=monthly_last_row + 1, column=1, value="Total").font = Font(bold=True)
ws_analysis.cell(row=monthly_last_row + 1, column=2, value=f"=SUM(B5:B{monthly_last_row})").font = Font(bold=True)
ws_analysis.cell(row=monthly_last_row + 1, column=3, value=f"=SUM(C5:C{monthly_last_row})").font = Font(bold=True)
ws_analysis.cell(row=monthly_last_row + 1, column=4, value=f"=SUM(D5:D{monthly_last_row})").font = Font(bold=True)
ws_analysis.cell(row=monthly_last_row + 1, column=5, value=f"=SUM(E5:E{monthly_last_row})").font = Font(bold=True)
ws_analysis.cell(row=monthly_last_row + 1, column=6, value=f"=SUM(F5:F{monthly_last_row})").font = Font(bold=True)
for col_idx in range(2, 7):
    ws_analysis.cell(row=monthly_last_row + 1, column=col_idx).number_format = "#,##0.00" if col_idx >= 4 else "#,##0"

p_start_row = monthly_last_row + 4
ws_analysis.cell(row=p_start_row - 1, column=1, value="Top 10 Products by Net Revenue").font = SECTION_FONT
prod_headers = ["Product Name", "Category", "Units Sold", "Gross Sales", "Discounts", "Net Revenue"]
for ci, h in enumerate(prod_headers, 1):
    c = ws_analysis.cell(row=p_start_row, column=ci, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")

top_10_products = product_sales.head(10).reset_index(drop=True)
for idx, r in top_10_products.iterrows():
    row_num = p_start_row + 1 + idx
    ws_analysis.cell(row=row_num, column=1, value=r["Product Name"])
    ws_analysis.cell(row=row_num, column=2, value=r["Category"])
    ws_analysis.cell(row=row_num, column=3, value=int(r["units_sold"])).number_format = "#,##0"
    ws_analysis.cell(row=row_num, column=4, value=float(r["gross_sales"])).number_format = "#,##0.00"
    ws_analysis.cell(row=row_num, column=5, value=float(r["discount_amount"])).number_format = "#,##0.00"
    ws_analysis.cell(row=row_num, column=6, value=float(r["revenue"])).number_format = "#,##0.00"

prod_last_row = p_start_row + len(top_10_products)

cat_start_row = 34
ws_analysis.cell(row=cat_start_row, column=1, value="Category Summary").font = SECTION_FONT
cat_headers = ["Category", "Units Sold", "Gross Sales", "Net Revenue"]
for ci, h in enumerate(cat_headers, 1):
    c = ws_analysis.cell(row=cat_start_row + 1, column=ci, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")

for idx, r in category_sales.iterrows():
    row_num = cat_start_row + 2 + idx
    ws_analysis.cell(row=row_num, column=1, value=r["Category"])
    ws_analysis.cell(row=row_num, column=2, value=int(r["units_sold"])).number_format = "#,##0"
    ws_analysis.cell(row=row_num, column=3, value=float(r["gross_sales"])).number_format = "#,##0.00"
    ws_analysis.cell(row=row_num, column=4, value=float(r["revenue"])).number_format = "#,##0.00"

cat_last_row = cat_start_row + 1 + len(category_sales)

loc_start_row = cat_last_row + 3
ws_analysis.cell(row=loc_start_row - 1, column=1, value="Store Location Summary").font = SECTION_FONT
loc_headers = ["Store Location", "Orders", "Units Sold", "Net Revenue"]
for ci, h in enumerate(loc_headers, 1):
    c = ws_analysis.cell(row=loc_start_row, column=ci, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")

for idx, r in location_sales.iterrows():
    row_num = loc_start_row + 1 + idx
    ws_analysis.cell(row=row_num, column=1, value=r["Store Location"])
    ws_analysis.cell(row=row_num, column=2, value=int(r["orders"])).number_format = "#,##0"
    ws_analysis.cell(row=row_num, column=3, value=int(r["units_sold"])).number_format = "#,##0"
    ws_analysis.cell(row=row_num, column=4, value=float(r["revenue"])).number_format = "#,##0.00"

loc_last_row = loc_start_row + len(location_sales)

ws_analysis.conditional_formatting.add(
    f"F5:F{monthly_last_row}",
    DataBarRule(start_type="min", end_type="max", color="176B87", showValue=True)
)
ws_analysis.conditional_formatting.add(
    f"F{p_start_row+1}:F{prod_last_row}",
    DataBarRule(start_type="min", end_type="max", color="64CCC5", showValue=True)
)

for ci in range(1, 8):
    ws_analysis.column_dimensions[get_column_letter(ci)].width = 20
ws_analysis.freeze_panes = "A4"

ws_dash.sheet_view.showGridLines = False

ws_dash.column_dimensions["A"].width = 3
for col_letter in ["B", "C", "D", "E", "F", "G"]:
    ws_dash.column_dimensions[col_letter].width = 13
ws_dash.column_dimensions["H"].width = 4
for col_letter in ["I", "J", "K", "L", "M", "N"]:
    ws_dash.column_dimensions[col_letter].width = 13
ws_dash.column_dimensions["O"].width = 3

ws_dash.row_dimensions[1].height = 12
ws_dash.row_dimensions[2].height = 30
ws_dash.row_dimensions[3].height = 16
ws_dash.row_dimensions[4].height = 8
ws_dash.row_dimensions[5].height = 16
ws_dash.row_dimensions[6].height = 26
ws_dash.row_dimensions[7].height = 4
for r_idx in range(8, 19):
    ws_dash.row_dimensions[r_idx].height = 16
ws_dash.row_dimensions[23].height = 8
for r_idx in range(24, 39):
    ws_dash.row_dimensions[r_idx].height = 16

ws_dash["B2"] = "Executive Sales Analysis Dashboard"
ws_dash["B2"].font = TITLE_FONT
start_dt = sales["Purchase Date"].min().date().isoformat()
end_dt = sales["Purchase Date"].max().date().isoformat()
ws_dash["B3"] = f"Period: {start_dt} to {end_dt}  |  Source: Sales_Data.xlsx  |  Amounts in workbook units; currency unspecified"
ws_dash["B3"].font = SUB_FONT

kpi_cards = [
    ("Gross Sales", f"=SUM('Cleaned Data'!I2:I{last_clean})", "#,##0.00", 2, 3, KPI_VALUE_FONT),
    ("Total Discounts", f"=SUM('Cleaned Data'!K2:K{last_clean})", "#,##0.00", 4, 5, KPI_VALUE_FONT),
    ("Net Revenue", f"=SUM('Cleaned Data'!L2:L{last_clean})", "#,##0.00", 6, 7, KPI_VALUE_FONT),
    ("Units Sold", f"=SUM('Cleaned Data'!G2:G{last_clean})", "#,##0", 9, 10, KPI_VALUE_FONT),
    ("Total Orders", f"=COUNTA('Cleaned Data'!A2:A{last_clean})", "#,##0", 11, 12, KPI_VALUE_FONT),
    ("Top Revenue Product", top_revenue_prod_name, "@", 13, 14, KPI_TEXT_FONT),
]

for label, formula, fmt, col_start, col_end, val_font in kpi_cards:
    ws_dash.merge_cells(start_row=5, start_column=col_start, end_row=5, end_column=col_end)
    top_cell = ws_dash.cell(row=5, column=col_start, value=label)
    top_cell.font = KPI_LABEL_FONT
    top_cell.fill = KPI_FILL
    top_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws_dash.merge_cells(start_row=6, start_column=col_start, end_row=6, end_column=col_end)
    val_cell = ws_dash.cell(row=6, column=col_start, value=formula)
    val_cell.font = val_font
    val_cell.fill = KPI_FILL
    val_cell.alignment = Alignment(horizontal="center", vertical="center")
    val_cell.number_format = fmt

    for r_idx in (5, 6):
        for c_idx in (col_start, col_end):
            ws_dash.cell(row=r_idx, column=c_idx).border = KPI_BOX_BORDER

chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Monthly Net Revenue Trend"
chart1.y_axis.title = "Net Revenue"
chart1.x_axis.tickLblSkip = 2
chart1.x_axis.tickMarkSkip = 2
chart1.x_axis.tickLblPos = "low"
chart1.height = 7.0
chart1.width = 13.5
data1 = Reference(ws_analysis, min_col=6, min_row=4, max_row=monthly_last_row)
cats1 = Reference(ws_analysis, min_col=1, min_row=5, max_row=monthly_last_row)
chart1.add_data(data1, titles_from_data=True)
chart1.set_categories(cats1)
chart1.legend = None
ws_dash.add_chart(chart1, "B8")

chart2 = BarChart()
chart2.type = "bar"
chart2.style = 11
chart2.title = "Top 10 Products by Net Revenue"
chart2.x_axis.title = "Net Revenue"
chart2.y_axis.title = "Product"
chart2.height = 7.0
chart2.width = 13.5
data2 = Reference(ws_analysis, min_col=6, min_row=p_start_row, max_row=prod_last_row)
cats2 = Reference(ws_analysis, min_col=1, min_row=p_start_row+1, max_row=prod_last_row)
chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)
chart2.legend = None
ws_dash.add_chart(chart2, "I8")

chart3 = BarChart()
chart3.type = "bar"
chart3.style = 12
chart3.title = "Net Revenue by Category"
chart3.x_axis.title = "Net Revenue"
chart3.y_axis.title = "Category"
chart3.x_axis.number_format = "#,##0"
chart3.height = 8.5
chart3.width = 13.5
data3 = Reference(ws_analysis, min_col=4, min_row=cat_start_row + 1, max_row=cat_last_row)
cats3 = Reference(ws_analysis, min_col=1, min_row=cat_start_row + 2, max_row=cat_last_row)
chart3.add_data(data3, titles_from_data=True)
chart3.set_categories(cats3)
chart3.legend = None
chart3.dataLabels = DataLabelList()
chart3.dataLabels.showVal = True
ws_dash.add_chart(chart3, "B22")

chart4 = BarChart()
chart4.type = "col"
chart4.style = 13
chart4.title = "Net Revenue by Store Location"
chart4.y_axis.title = "Net Revenue"
chart4.x_axis.title = "Store Location"
chart4.height = 8.5
chart4.width = 13.5
data4 = Reference(ws_analysis, min_col=4, min_row=loc_start_row, max_row=loc_last_row)
cats4 = Reference(ws_analysis, min_col=1, min_row=loc_start_row+1, max_row=loc_last_row)
chart4.add_data(data4, titles_from_data=True)
chart4.set_categories(cats4)
chart4.legend = None
ws_dash.add_chart(chart4, "I22")

ws_dash.sheet_properties.pageSetUpPr.fitToPage = True
ws_dash.page_setup.orientation = "landscape"
ws_dash.page_setup.paperSize = ws_dash.PAPERSIZE_A4
ws_dash.page_setup.fitToWidth = 1
ws_dash.page_setup.fitToHeight = 1
ws_dash.print_area = "A1:O39"
ws_dash.sheet_view.zoomScale = 65
ws_dash.page_margins.left = 0.2
ws_dash.page_margins.right = 0.2
ws_dash.page_margins.top = 0.25
ws_dash.page_margins.bottom = 0.25
ws_dash.page_margins.header = 0.1
ws_dash.page_margins.footer = 0.1
ws_raw.print_area = "A1:L1479"
ws_clean.print_area = "A1:L1479"
ws_analysis.print_area = "A1:F49"

for ws in [ws_dash, ws_raw, ws_clean, ws_analysis]:
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4

wb.active = 0
wb.save(OUT_FILE)
print(f"Generated: {OUT_FILE}")
